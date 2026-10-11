"""Text extraction for uploads: PDFs with PyMuPDF, scans and photos with PaddleOCR, Word files.

Each PDF page is inspected before it is read:

  text   a real text layer: read it directly (exact characters, no confidence needed)
  ocr    no usable text layer: render the page and OCR it. This covers plain scans, scans that
         carry an invisible text layer from a scanner's own OCR (we do not trust its dates), pages
         whose text layer is garbage (broken font encodings), and scans with only a thin typed stamp
         on top, such as an e-filing or digital-signature footer
  mixed  a typed page that also embeds scanned images (a pasted stamp, an annexure snippet): the
         text layer is read directly and each image region is OCR'd, in reading order

Every page keeps its text, the method used and, for OCR'd text, the confidence of each word as
character spans. Date detection uses those spans to decide whether a date was read clearly enough
to be trusted, or must be quarantined.
"""
from __future__ import annotations

import logging
import re
from dataclasses import dataclass, field
from pathlib import Path

import numpy as np
from PIL import Image

from ..config import settings
from . import ocr

log = logging.getLogger(__name__)

IMAGE_EXT = {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".webp", ".bmp", ".heic", ".heif"}
PDF_EXT = {".pdf"}
WORD_EXT = {".docx"}
ALLOWED_EXT = IMAGE_EXT | PDF_EXT | WORD_EXT


@dataclass
class PageText:
    page_no: int
    text: str
    method: str  # text | ocr | mixed
    confidence: float | None = None  # mean OCR word confidence (0-100) of the OCR'd words
    words: list[list[float]] = field(default_factory=list)  # [start, end, conf] for OCR'd words only


class ExtractionError(Exception):
    pass


def _ocr(img, page_no: int) -> PageText:
    try:
        return ocr.ocr_image(img, page_no)
    except ocr.OCRUnavailable as e:
        raise ExtractionError(f"This file needs OCR, but the OCR engine is unavailable. {e}") from e


# --------------------------------------------------------------------------- PDF


def _scan_dpi(page, clip=None) -> int:
    """Resolution to render at: the scan's own resolution (no point upsampling a 200 dpi scan),
    between 150 dpi and OCR_DPI. Pages with no embedded image use OCR_DPI."""
    import pymupdf

    best = 0.0
    area = clip or page.rect
    for info in page.get_image_info():
        r = pymupdf.Rect(info["bbox"]) & area
        if r.is_empty or r.width < 36:
            continue
        full = pymupdf.Rect(info["bbox"])
        best = max(best, info.get("width", 0) / (full.width / 72.0))
    if best <= 0:
        return settings.ocr_dpi
    return int(min(settings.ocr_dpi, max(150.0, best)))


def _render(page, clip=None, blank=()) -> np.ndarray:
    """Render a page (or part of it) to a BGR array for OCR. Rectangles in `blank` (text-layer
    words) are painted white so OCR only sees what the text layer does not already hold."""
    import pymupdf

    dpi = _scan_dpi(page, clip)
    pix = page.get_pixmap(dpi=dpi, clip=clip, alpha=False, colorspace=pymupdf.csRGB)
    rgb = np.frombuffer(pix.samples, dtype=np.uint8).reshape(pix.height, pix.width, pix.n)[:, :, :3]
    bgr = np.ascontiguousarray(rgb[:, :, ::-1])
    if blank:
        origin = clip or page.rect
        k = dpi / 72.0
        for r in blank:
            x0, y0 = int((r.x0 - origin.x0) * k) - 2, int((r.y0 - origin.y0) * k) - 2
            x1, y1 = int((r.x1 - origin.x0) * k) + 3, int((r.y1 - origin.y0) * k) + 3
            bgr[max(0, y0):max(0, y1), max(0, x0):max(0, x1)] = 255
    return bgr


_PLAIN = set(".,;:'\"()[]{}-/&%@#!?+*=_<>|~`$₹§–—’‘“”…°")


def _garbled(text: str) -> bool:
    """A text layer that exists but is unusable: replacement characters, '(cid:12)' runs from fonts
    without a Unicode map, or 'words' with no vowels from a broken encoding."""
    if not text.strip():
        return True
    bad = text.count("�") + 5 * len(re.findall(r"\(cid:\d+\)", text))
    plain = sum(1 for c in text if c.isalnum() or c.isspace() or c in _PLAIN)
    if (plain - bad) / len(text) < 0.85:
        return True
    words = re.findall(r"\b[A-Za-z]{3,}\b", text)
    if len(words) >= 8:
        voweled = sum(1 for w in words if re.search(r"[aeiouyAEIOUY]", w))
        if voweled / len(words) < 0.5:
            return True
    return False


def _coverage(rects: list, page_rect, grid: int = 48) -> float:
    """Share of the page covered by the given rectangles (overlaps counted once)."""
    if not rects:
        return 0.0
    W, H = page_rect.width or 1, page_rect.height or 1
    occ = np.zeros((grid, grid), dtype=bool)
    for r in rects:
        x0, x1 = int(max(0, (r.x0 - page_rect.x0) / W * grid)), int(min(grid, np.ceil((r.x1 - page_rect.x0) / W * grid)))
        y0, y1 = int(max(0, (r.y0 - page_rect.y0) / H * grid)), int(min(grid, np.ceil((r.y1 - page_rect.y0) / H * grid)))
        if x1 > x0 and y1 > y0:
            occ[y0:y1, x0:x1] = True
    return float(occ.mean())


def _image_rects(page) -> list:
    import pymupdf

    out = []
    for info in page.get_image_info():
        r = pymupdf.Rect(info["bbox"]) & page.rect
        if not r.is_empty and r.width > 1 and r.height > 1:
            out.append(r)
    return out


def _plan(page) -> str:
    """Decide how to read one PDF page: 'text', 'ocr' or 'mixed'."""
    visible = invisible = 0
    try:
        for span in page.get_texttrace():
            n = len(span.get("chars", ()))
            if span.get("type") == 3 or span.get("opacity", 1) == 0:
                invisible += n
            else:
                visible += n
    except Exception:  # noqa: BLE001 - very old PyMuPDF or odd content streams: fall back to plain text
        visible = len(page.get_text("text"))
    text = page.get_text("text") or ""
    images = _image_rects(page)
    img_cov = _coverage(images, page.rect)

    if visible < settings.text_layer_min_chars:
        return "ocr"  # no text layer, or only a scanner's invisible OCR layer
    if _garbled(text):
        return "ocr"
    if img_cov >= 0.5:
        # a scan with a typed stamp on top (e-filing footer, digital signature box)? then the
        # words cover little of the page and the scan underneath holds the real content
        words = [w for w in page.get_text("words")]
        import pymupdf

        word_cov = _coverage([pymupdf.Rect(w[:4]) for w in words], page.rect)
        if word_cov < 0.10:
            return "ocr"
    big = [r for r in images if r.width * r.height >= 0.02 * page.rect.width * page.rect.height
           and min(r.width, r.height) >= 28]
    return "mixed" if big else "text"


def _tidy(text: str) -> str:
    lines = [ln.rstrip() for ln in text.replace("\r", "").split("\n")]
    out = "\n".join(lines)
    while "\n\n\n" in out:
        out = out.replace("\n\n\n", "\n\n")
    return out.strip()


def _text_blocks(page) -> list[tuple[float, float, str]]:
    """Text-layer blocks as (top, left, text), in reading order."""
    out = []
    for b in page.get_text("blocks", sort=True):
        if b[6] != 0:  # image block
            continue
        t = _tidy(b[4])
        if t:
            out.append((b[1], b[0], t))
    return out


def _mixed_page(page, page_no: int) -> PageText:
    """Typed page with embedded scans: text layer as is, plus OCR of each sizeable image region."""
    import pymupdf

    blocks: list[tuple[float, float, str, PageText | None]] = [(y, x, t, None) for y, x, t in _text_blocks(page)]
    words_on_page = [pymupdf.Rect(w[:4]) for w in page.get_text("words")]
    area = page.rect.width * page.rect.height
    regions = sorted((r for r in _image_rects(page) if r.width * r.height >= 0.02 * area and min(r.width, r.height) >= 28),
                     key=lambda r: -r.width * r.height)[:6]
    for r in regions:
        # typed words drawn over the image are already in the text layer: blank them out first
        img = _render(page, clip=r, blank=[w for w in words_on_page if w.intersects(r)])
        if (img.min(axis=2) < 160).mean() < 0.002:
            continue  # nothing left but paper
        pt = _ocr(img, page_no)
        if any(pt.text[int(w[0]):int(w[1])] != ocr.ILLEGIBLE for w in pt.words):
            blocks.append((r.y0, r.x0, pt.text, pt))
    blocks.sort(key=lambda b: (round(b[0] / 6), b[1]))

    parts, words, confs, pos = [], [], [], 0
    for _, _, text, pt in blocks:
        if parts:
            parts.append("\n\n")
            pos += 2
        parts.append(text)
        if pt is not None:
            words.extend([[w[0] + pos, w[1] + pos, w[2]] for w in pt.words])
            confs.extend(w[2] for w in pt.words if pt.text[int(w[0]):int(w[1])] != ocr.ILLEGIBLE)
        pos += len(text)
    if not words:
        return PageText(page_no, "".join(parts), "text")
    return PageText(page_no, "".join(parts), "mixed", round(sum(confs) / len(confs), 1) if confs else 0.0, words)


def _from_pdf(path: Path) -> tuple[list[PageText], str]:
    import pymupdf

    try:
        doc = pymupdf.open(path)
    except Exception as e:  # noqa: BLE001
        raise ExtractionError(f"Could not open the PDF: {e}") from e
    if doc.needs_pass and not doc.authenticate(""):
        raise ExtractionError("The PDF is password-protected. Upload an unlocked copy.")
    if doc.page_count == 0:
        raise ExtractionError("The PDF has no pages.")
    pages: list[PageText] = []
    with doc:
        for i, page in enumerate(doc, start=1):
            plan = _plan(page)
            if plan == "ocr":
                pages.append(_ocr(_render(page), i))
            elif plan == "mixed":
                pages.append(_mixed_page(page, i))
            else:
                text = "\n\n".join(t for _, _, t in _text_blocks(page))
                pages.append(PageText(page_no=i, text=text, method="text"))
            log.debug("page %s read as %s", i, pages[-1].method)
    methods = {p.method for p in pages}
    label = "Scanned PDF" if methods == {"ocr"} else "PDF"
    return pages, label


# --------------------------------------------------------------------------- images and Word


def _from_image(path: Path) -> tuple[list[PageText], str]:
    if path.suffix.lower() in {".heic", ".heif"}:
        try:
            import pillow_heif  # type: ignore

            pillow_heif.register_heif_opener()
        except ImportError as e:
            raise ExtractionError("HEIC images need the pillow-heif package on the server.") from e
    try:
        img = Image.open(path)
        img.load()
    except Exception as e:  # noqa: BLE001
        raise ExtractionError(f"Could not open the image: {e}") from e
    pages = []
    for i in range(getattr(img, "n_frames", 1)):  # multi-page TIFFs
        img.seek(i)
        pages.append(_ocr(img.copy(), i + 1))
    return pages, "Scanned image"


def _from_docx(path: Path) -> tuple[list[PageText], str]:
    import docx

    try:
        d = docx.Document(str(path))
    except Exception as e:  # noqa: BLE001
        raise ExtractionError(f"Could not open the Word file: {e}") from e
    paras = [p.text for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            paras.append("  ".join(c.text for c in row.cells))
    text = "\n\n".join(p for p in paras if p.strip())
    # Word has no fixed pages; split roughly every 3,000 characters on paragraph breaks
    pages, buf, n = [], [], 1
    for p in text.split("\n\n"):
        buf.append(p)
        if sum(len(x) for x in buf) > 3000:
            pages.append(PageText(n, "\n\n".join(buf), "text"))
            buf, n = [], n + 1
    if buf or not pages:
        pages.append(PageText(n, "\n\n".join(buf), "text"))
    return pages, "Word"


def extract_pages(path: Path) -> tuple[list[PageText], str]:
    """Return (pages, format label) for a stored upload."""
    ext = path.suffix.lower()
    if ext in PDF_EXT:
        return _from_pdf(path)
    if ext in IMAGE_EXT:
        return _from_image(path)
    if ext in WORD_EXT:
        return _from_docx(path)
    raise ExtractionError(f"Unsupported file type {ext}. Upload a PDF, a scanned image or a Word file.")


def span_confidence(words: list[list[float]], start: int, end: int) -> float | None:
    """Lowest OCR confidence of the words overlapping [start, end); None for text-layer text."""
    hits = [w[2] for w in words if w[0] < end and w[1] > start]
    return min(hits) if hits else None
