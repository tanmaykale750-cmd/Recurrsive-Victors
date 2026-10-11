"""Court-ready PDF (ReportLab) and Word (python-docx) exports of the chronology."""
from __future__ import annotations

import io

NAVY_HEX = "1F3A5F"


# --------------------------------------------------------------------------- PDF

def to_pdf(r: dict) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import LongTable, Paragraph, SimpleDocTemplate, Spacer, TableStyle

    navy = colors.HexColor("#" + NAVY_HEX)
    st = {
        "h": ParagraphStyle("h", fontName="Times-Bold", fontSize=15, leading=19, alignment=1, spaceAfter=6),
        "c": ParagraphStyle("c", fontName="Times-Roman", fontSize=11.5, leading=15, alignment=1),
        "ci": ParagraphStyle("ci", fontName="Times-Italic", fontSize=11.5, leading=15, alignment=1, spaceAfter=8),
        "small": ParagraphStyle("s", fontName="Helvetica", fontSize=8.5, leading=11, alignment=1, textColor=colors.HexColor("#555555")),
        "sec": ParagraphStyle("sec", fontName="Times-Bold", fontSize=13, leading=16, textColor=navy, spaceBefore=14, spaceAfter=3),
        "note": ParagraphStyle("n", fontName="Helvetica", fontSize=8.5, leading=11, textColor=colors.HexColor("#555555"), spaceAfter=6),
        "cell": ParagraphStyle("cell", fontName="Times-Roman", fontSize=9, leading=11.5),
        "flag": ParagraphStyle("flag", fontName="Helvetica-Bold", fontSize=8.5, leading=11, textColor=colors.HexColor("#B42318")),
        "th": ParagraphStyle("th", fontName="Helvetica-Bold", fontSize=8.5, leading=11, textColor=colors.white),
    }
    W = A4[0] - 36 * mm

    def esc(s) -> str:
        return str(s if s is not None else "").replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")

    def table(cols, rows, flag_key=None):
        data = [[Paragraph(esc(c["label"]), st["th"]) for c in cols]]
        for row in rows:
            data.append([Paragraph(esc(row.get(c["key"], "")), st["flag"] if flag_key == c["key"] and row.get(c["key"]) else st["cell"]) for c in cols])
        t = LongTable(data, colWidths=[c["w"] * W for c in cols], repeatRows=1)
        t.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), navy), ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#C9CED8")),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#FAFBFC")]),
            ("LEFTPADDING", (0, 0), (-1, -1), 5), ("RIGHTPADDING", (0, 0), (-1, -1), 5), ("TOPPADDING", (0, 0), (-1, -1), 4), ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ]))
        return t

    m, cols = r["meta"], r["cols"]
    story = [Paragraph(esc(m["heading"]), st["h"]), Paragraph(esc(m["matter"]), st["c"]), Paragraph(esc(m["parties"]), st["ci"])]
    if m.get("court"):
        story.append(Paragraph(esc(m["court"]), st["c"]))
    story += [Spacer(1, 4), Paragraph(esc(m["prepared"]), st["small"]), Spacer(1, 6)]
    story += [Paragraph("Part A: Verified chronology", st["sec"]),
              Paragraph("Dates stated in the record, or confirmed by the reviewer where marked. Each entry is tied to its source.", st["note"]),
              table(cols["verified"], r["verified"])]
    if r["unresolved"]:
        story += [Paragraph("Part B: Unresolved issues", st["sec"]),
                  Paragraph("Not verified. Listed separately so the record can be checked; no date has been chosen or calculated.", st["note"]),
                  table(cols["unresolved"], r["unresolved"], "type")]
    if r["missing"]:
        story += [Paragraph("Appendix A: Documents referred to but not in the bundle", st["sec"]), table(cols["missing"], r["missing"])]
    if r["corrections"]:
        story += [Paragraph("Appendix B: Reviewer correction log", st["sec"]), table(cols["log"], r["corrections"])]

    footer = f"CaseCompass chronology. {m['disclaimer']}"

    def on_page(canvas, doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 7.5)
        canvas.setFillColor(colors.HexColor("#666666"))
        canvas.drawString(18 * mm, 10 * mm, footer[:150])
        canvas.drawRightString(A4[0] - 18 * mm, 10 * mm, f"Page {doc.page}")
        canvas.restoreState()

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm, topMargin=18 * mm, bottomMargin=18 * mm,
                            title="Chronology of events and list of dates", author="CaseCompass")
    doc.build(story, onFirstPage=on_page, onLaterPages=on_page)
    return buf.getvalue()


# --------------------------------------------------------------------------- Word

def to_docx(r: dict) -> bytes:
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    d = Document()
    sec = d.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, side, Cm(1.8))
    normal = d.styles["Normal"]
    normal.font.name, normal.font.size = "Times New Roman", Pt(10.5)

    def para(text, *, bold=False, italic=False, size=10.5, center=False, color=None, sans=False, after=4, before=0):
        p = d.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_after, p.paragraph_format.space_before = Pt(after), Pt(before)
        run = p.add_run(text)
        run.bold, run.italic, run.font.size = bold, italic, Pt(size)
        run.font.name = "Arial" if sans else "Times New Roman"
        if color:
            run.font.color.rgb = RGBColor.from_string(color)
        return p

    def shade(cell, hex_fill):
        tcPr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), hex_fill)
        tcPr.append(shd)

    def table(cols, rows, flag_key=None):
        t = d.add_table(rows=1, cols=len(cols))
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        total = 17.4
        for i, c in enumerate(cols):
            cell = t.rows[0].cells[i]
            cell.width = Cm(total * c["w"])
            shade(cell, NAVY_HEX)
            run = cell.paragraphs[0].add_run(c["label"])
            run.bold, run.font.size, run.font.name = True, Pt(8.5), "Arial"
            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        # repeat header row on each page
        trPr = t.rows[0]._tr.get_or_add_trPr()
        el = OxmlElement("w:tblHeader")
        el.set(qn("w:val"), "true")
        trPr.append(el)
        for row in rows:
            cells = t.add_row().cells
            for i, c in enumerate(cols):
                cells[i].width = Cm(total * c["w"])
                v = str(row.get(c["key"], "") or "")
                run = cells[i].paragraphs[0].add_run(v)
                run.font.size = Pt(9)
                if flag_key == c["key"] and v:
                    run.bold, run.font.name = True, "Arial"
                    run.font.color.rgb = RGBColor(0xB4, 0x23, 0x18)
        d.add_paragraph().paragraph_format.space_after = Pt(2)

    m, cols = r["meta"], r["cols"]
    para(m["heading"], bold=True, size=14, center=True, after=6)
    para(m["matter"], size=11.5, center=True, after=2)
    para(m["parties"], italic=True, size=11.5, center=True, after=2)
    if m.get("court"):
        para(m["court"], size=11, center=True, after=6)
    para(m["prepared"], size=8.5, center=True, sans=True, color="5A5F6B", after=10)
    para("Part A: Verified chronology", bold=True, size=12.5, color=NAVY_HEX, before=6, after=2)
    para("Dates stated in the record, or confirmed by the reviewer where marked.", size=8.5, sans=True, color="5A5F6B", after=6)
    table(cols["verified"], r["verified"])
    if r["unresolved"]:
        para("Part B: Unresolved issues", bold=True, size=12.5, color=NAVY_HEX, before=10, after=2)
        para("Not verified. No date has been chosen or calculated for these entries.", size=8.5, sans=True, color="5A5F6B", after=6)
        table(cols["unresolved"], r["unresolved"], "type")
    if r["missing"]:
        para("Appendix A: Documents referred to but not in the bundle", bold=True, size=12.5, color=NAVY_HEX, before=10, after=6)
        table(cols["missing"], r["missing"])
    if r["corrections"]:
        para("Appendix B: Reviewer correction log", bold=True, size=12.5, color=NAVY_HEX, before=10, after=6)
        table(cols["log"], r["corrections"])

    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = fp.add_run("CaseCompass chronology. Organises the record; not legal advice. Page ")
    run.font.size = Pt(7.5)
    for tag, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if tag:
            fc = OxmlElement("w:fldChar")
            fc.set(qn("w:fldCharType"), tag)
            r_ = fp.add_run()
            r_._r.append(fc)
        else:
            it = OxmlElement("w:instrText")
            it.set(qn("xml:space"), "preserve")
            it.text = text
            r_ = fp.add_run()
            r_._r.append(it)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()
